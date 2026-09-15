"""Read existing company rules once per run; retain the last valid snapshot."""
from __future__ import annotations

import json
import logging
import os
from pathlib import Path
from .models import utc_now_iso, normalized_text
from .rir_products import usable_name
from common.keyword_watch import normalize_keyword_term


SHEETS = {'rs': 'pc_palabras_clave', 'negative': 'pc_palabras_negativas', 'fichas': 'ct_rir_fichas'}


def read_existing_profiles():
    from google.oauth2.service_account import Credentials
    from orquestador.google_transport import build_sheets_service, close_service, retry_google_call
    credential_path = os.getenv('ORQUESTADOR_PANAMACOMPRA_SERVICE_ACCOUNT_FILE') or str(Path(__file__).resolve().parents[1] / 'credentials/service-account.json')
    creds = Credentials.from_service_account_file(credential_path, scopes=['https://www.googleapis.com/auth/spreadsheets.readonly'])
    client = None
    def reset():
        nonlocal client
        close_service(client)
        client = None
    def call():
        nonlocal client
        client = client or build_sheets_service(creds)
        return client.spreadsheets().values().batchGet(
            spreadsheetId=os.getenv('ORQUESTADOR_PANAMACOMPRA_SPREADSHEET_ID', '17hOfP-vMdJ4D7xym1cUp7vAcd8XJPErpY3V-9Ui2tCo'),
            ranges=[f"'{sheet}'!A1:B" for sheet in SHEETS.values()]).execute()
    try:
        result = retry_google_call(call, reset=reset, label='perfiles de oportunidades externas')
        return dict(zip(SHEETS.values(), (r.get('values', []) for r in result.get('valueRanges', []))))
    finally:
        reset()


def read_catalog_names(codes):
    """Resolve missing watchlist names from the catalogs already on the server."""
    from openpyxl import load_workbook
    root = Path(__file__).resolve().parents[1]
    candidates = [Path(p).expanduser() for p in os.getenv('FICHAS_CATALOG_PATHS', '').split(os.pathsep) if p.strip()]
    candidates.extend([root/'data/fichas/fichas-y-nombre.xlsx',
                       Path.home()/'fichas/fichas-y-nombre.xlsx',
                       Path.home()/'GEAPP/fichas_ctni_con_enlace.xlsx'])
    names = {}
    for path in dict.fromkeys(candidates):
        if not path.exists() or path.suffix.lower() != '.xlsx':
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            # Existing two-column catalogs can be headerless; read the first row too.
            for row in workbook.active.iter_rows(min_col=1, max_col=2, values_only=True):
                code = str(row[0] or '').strip().split('.')[0]
                name = usable_name(row[1])
                if code in codes and code not in names and name:
                    names[code] = name
                    if len(names) == len(codes):
                        return names
        finally:
            workbook.close()
    return names


def load_profiles(path=None, reader=None, name_resolver=None):
    path = Path(path or os.getenv('OTRAS_FUENTES_PROFILE_PATH') or Path(__file__).resolve().parents[1] / 'data/otras_fuentes/profiles.json')
    try:
        previous = json.loads(path.read_text(encoding='utf8')) if path.exists() else {}
    except (OSError, ValueError):
        previous = {}
    try:
        if reader is None:
            batches = read_existing_profiles()
            reader = lambda sheet: batches.get(sheet, [])
        result = {}
        for key, sheet in SHEETS.items():
            rows = reader(sheet)
            if not rows:
                raise ValueError('Hoja sin encabezado; se conserva el perfil anterior')
            values = [str(row[0]).strip() for row in rows if row]
            if key == 'fichas':
                result[key] = sorted({v.strip('* ').split('.')[0] for v in values if v.strip('* ').split('.')[0].isdigit()})
                names = {str(row[0]).strip('* ').split('.')[0]: usable_name(row[1])
                         for row in rows if len(row) > 1}
                cached = {p['ficha']: p for p in previous.get('rir_products', [])}
                missing = {code for code in result[key] if not names.get(code)}
                resolved = {}
                if missing:
                    try:
                        resolved = (name_resolver or read_catalog_names)(missing)
                    except Exception as exc:
                        logging.getLogger('otras_fuentes').warning('Nombres RIR: %s; conservando nombres anteriores', type(exc).__name__)
                products = []
                unresolved = []
                for code in result[key]:
                    name = names.get(code) or usable_name(resolved.get(code)) or usable_name(cached.get(code, {}).get('name'))
                    if name:
                        origin = 'ct_rir_fichas' if names.get(code) else ('catalogo_existente' if resolved.get(code) else 'cache_anterior')
                        products.append({'ficha': code, 'name': name, 'source': origin})
                    else:
                        unresolved.append(code)
                result['rir_products'] = products
                result['rir_unresolved_fichas'] = unresolved
            else:
                result[key] = list(dict.fromkeys(normalize_keyword_term(v) for v in values if normalized_text(v) not in {'palabra clave', 'palabras clave', 'keyword'} and normalize_keyword_term(v)))
        result['loaded_at'] = utc_now_iso()
        result['rir_basis'] = 'product_words_v1'
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix('.tmp')
        temporary.write_text(json.dumps(result, ensure_ascii=False), encoding='utf8')
        temporary.replace(path)
        return {**result, 'status': 'updated'}
    except Exception as exc:
        logging.getLogger('otras_fuentes').warning('Perfil de empresas: %s; usando último perfil válido o reglas iniciales', type(exc).__name__)
        return {**previous, 'status': 'cached' if previous else 'defaults'}
